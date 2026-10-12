"""Shared test fixtures. Builds a small FAKE CV — never use the real one in tests."""

from pathlib import Path

import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def _add_hyperlink(paragraph, text: str, url: str) -> None:
    """python-docx has no add_hyperlink(), so build the XML by hand."""
    r_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


@pytest.fixture
def fake_cv_docx(tmp_path: Path) -> Path:
    doc = Document()
    doc.add_paragraph("Jane Doe", style="Title")

    doc.add_heading("Work History", level=1)
    doc.add_paragraph("Data Analyst | Acme Ltd | London, UK | 01/2024 – 06/2025")
    doc.add_paragraph("Built a Python ETL pipeline with Airflow.", style="List Bullet")
    doc.add_paragraph("•  Ran 10+ A/B tests on pricing.")  # typed bullet, no Word list

    doc.add_heading("Core Skills", level=1)
    doc.add_heading("Programming Languages", level=3)
    doc.add_paragraph("Python, SQL")

    doc.add_heading("Projects", level=1)
    p = doc.add_paragraph("Demo Project | Python | ")
    _add_hyperlink(p, "GitHub", "https://github.com/jane/demo")
    p = doc.add_paragraph("Broken Link Project | Python | ")
    _add_hyperlink(p, "GitHub", "○%09https:/github.com/jane/broken")  # copy-paste damage

    path = tmp_path / "fake_cv.docx"
    doc.save(path)
    return path
